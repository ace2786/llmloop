#!/usr/bin/env python3
"""判分辅助工具：把待判条目导出成单页 HTML，判完导出 CSV 再回填 Excel。

★ 这个工具不代替人判分。它只做三件事：
  1. 把 82 条待判条目摊成一屏一条，省掉 Excel 里上下滚的时间
  2. 键盘 1/2/3 判分，进度自动存浏览器，关掉不丢
  3. 导出 CSV，再由 apply 子命令写回 Excel 的 F/G/H 三列

判分仍然由人做，RUBRIC.md 的盲判规则不变（HTML 里没有任何版本/模型信息）。

用法：
    python scripts/grade_aid.py make                    # 生成 out/grade_v3.html
    python scripts/grade_aid.py apply <导出的csv>        # 写回 Excel
"""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "out"
DEFAULT_XLSX = Path("/mnt/d/llmloop/待判分_v3的120条.xlsx")

COLS = ["代号", "题号", "问题", "回答", "检索到的段落", "人工判分", "失败原因", "判断理由", "类型", "标准答案"]


def read_rows(xlsx: Path) -> list[dict]:
    wb = load_workbook(xlsx, data_only=True)
    ws = wb["判分"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r or r[0] is None:
            continue
        d = {k: ("" if v is None else str(v)) for k, v in zip(COLS, r)}
        rows.append(d)
    return rows


def do_make(xlsx: Path) -> int:
    rows = read_rows(xlsx)
    todo = [r for r in rows if not r["人工判分"].strip()]
    print(f"总计 {len(rows)} 条，已判 {len(rows) - len(todo)} 条，待判 {len(todo)} 条")

    items = [
        {
            "code": r["代号"],
            "qid": r["题号"],
            "type": r["类型"],
            "q": r["问题"],
            "ctx": r["检索到的段落"],
            "ref": r["标准答案"],
            "ans": r["回答"],
        }
        for r in todo
    ]

    tpl = (Path(__file__).resolve().parent / "grade_aid_template.html").read_text(encoding="utf-8")
    html = tpl.replace("/*__DATA__*/", json.dumps(items, ensure_ascii=False))
    dst = OUT / "grade_v3.html"
    dst.write_text(html, encoding="utf-8")
    print(f"已生成：{dst}")
    print(f"Windows 路径：\\\\wsl.localhost\\Ubuntu\\home\\ace\\llmloop\\out\\grade_v3.html")
    print()
    print("  快捷键：1=对  2=部分对  3=错  空格=跳过  ←=上一条")
    print("  判「错/部分对」时再按 1-4 选失败原因，回车进下一条")
    return 0


def do_apply(xlsx: Path, csv_path: Path) -> int:
    with csv_path.open(encoding="utf-8-sig") as f:
        grades = {r["代号"].strip(): r for r in csv.DictReader(f)}
    print(f"读到判分 {len(grades)} 条")

    backup = xlsx.with_suffix(".bak.xlsx")
    shutil.copy2(xlsx, backup)
    print(f"已备份原表：{backup}")

    wb = load_workbook(xlsx)
    ws = wb["判分"]
    n = 0
    for i in range(2, ws.max_row + 1):
        code = ws.cell(row=i, column=1).value
        if not code:
            continue
        g = grades.get(str(code).strip())
        if not g:
            continue
        ws.cell(row=i, column=6, value=(g.get("人工判分") or "").strip())
        ws.cell(row=i, column=7, value=(g.get("失败原因") or "").strip() or None)
        ws.cell(row=i, column=8, value=(g.get("判断理由") or "").strip() or None)
        n += 1
    wb.save(xlsx)
    print(f"已写回 {n} 条到 {xlsx}")
    print()
    print("下一步：python scripts/blind_xlsx.py merge")
    return 0


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["make", "apply"])
    ap.add_argument("csv", nargs="?", help="apply 时要写入的 CSV")
    ap.add_argument("--xlsx", default=str(DEFAULT_XLSX))
    a = ap.parse_args()
    xlsx = Path(a.xlsx)
    if not xlsx.exists():
        print(f"找不到判分表：{xlsx}", file=sys.stderr)
        return 1
    if a.cmd == "make":
        return do_make(xlsx)
    if not a.csv:
        print("apply 需要指定 CSV 路径", file=sys.stderr)
        return 1
    return do_apply(xlsx, Path(a.csv))


if __name__ == "__main__":
    sys.exit(main())
